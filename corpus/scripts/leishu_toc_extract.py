#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
leishu_toc_extract.py  (v2)

Recover leaf-level category trees (原目) from leishu full texts already held
in the Fanya Hanwen corpus, for the Browse-by-Category workstream.

Design rules
------------
1. Nothing is invented. A label is emitted only if it is literally present in
   the transcribed text.
2. Container (部/類/門/篇…) status is decided by a CLOSED VOCABULARY read from
   the attested top-level labels already recorded in the workbook sheet
   類書分類原目 — never by guessing from a terminal character. Suffix guessing
   promotes 百部 (a plant), 蒲類 (a polity) and 玉篇 (a book) to categories.
3. Parents are carried forward from the last attested container in reading
   order. Every row records whether that container was attested in the same
   juan or carried across a juan boundary, so nothing implies more structure
   than the witness supports.
4. Works whose text yields ENTRIES rather than categories are labelled at
   level 條目 so they can never be mistaken for a category tree.

Output: UTF-8 with BOM CSV, schema identical to 類書分類原目.
"""
import os, re, csv, json, glob

HOME    = os.path.expanduser('~')
CORPUS  = os.path.join(HOME, 'mnt/fanyahanwen-corpus/corpus')
BOOK    = os.path.join(HOME, 'mnt/fanyahanwen-corpus/fanya_category_grouping_2026-09-11.xlsx')
OUTDIR  = os.path.join(HOME, 'leishu')

BAD   = set('．。，，、：；:!？?「」『』（）()〈〉《》【】[]　 ／/…·•←→◄►●○※◆0123456789')
NOISE = {'返回頁首','姊妹计划','数据项','目錄','目录','无序列表项','無序列表項','校記','校记','編輯','编辑'}
FRONT = ('序','跋','凡例','小引','總序','总序','題辭','题辞','自序','後序','进书表','進書表')
# edition-provenance strings that look like headings in 四庫全書總目提要
PROV  = re.compile(r'(藏本|採進本|采进本|刊本|寫本|写本|抄本|大典本|通行本|進本|进本)$')

def load_vocab():
    """Closed container vocabulary, per work, from 類書分類原目."""
    import openpyxl
    wb = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
    ws = wb['類書分類原目']
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    ix = {h: i for i, h in enumerate(hdr)}
    CONT = {'部','類','門','篇','典','彙編','大類','門類','部門','存世類','卷目','略','集'}
    voc = {}
    for r in it:
        if not r or not r[ix['Work']]: continue
        w   = str(r[ix['Work']]).strip()
        lab = str(r[ix['Exact source label']] or '').strip()
        par = str(r[ix['Parent']] or '').strip()
        lvl = str(r[ix['Source level']] or '').strip()
        s = voc.setdefault(w, set())
        if par: s.add(par)
        if lab and lvl in CONT and not par: s.add(lab)
    # 四庫全書總目提要 inherits the 四部 containers recorded on the 四庫全書 sheet
    if '四庫全書' in wb.sheetnames:
        s = voc.setdefault('四庫全書總目提要', set())
        for r in wb['四庫全書'].iter_rows(min_row=2, values_only=True):
            if r and r[0]: s.add(str(r[0]).strip())
    return {k: expand(v) for k, v in voc.items() if v}

# The workbook records bare stems (帝王, 天); the transcriptions carry the
# taxonomy-marked forms (帝王部, 天部上). Accept the marked variants of an
# ATTESTED stem only. No new stems are ever created.
MARKS = ('', '部', '類', '門', '篇', '典')
POS   = ('', '上', '中', '下')
def expand(labels):
    out = set()
    for lab in labels:
        stem = lab
        for m in ('部','類','門','篇','典'):
            if stem.endswith(m) and len(stem) > 1: stem = stem[:-1]; break
        if stem and stem[-1] in POS[1:] and len(stem) > 1: stem = stem[:-1]
        for m in MARKS:
            for p in POS:
                out.add(stem + m + p)
        out.add(lab)
    return {x for x in out if x}

def read(p):
    with open(p, encoding='utf-8-sig', errors='replace') as f: return f.read()

def meta_of(text):
    m = {}
    for line in text.split('\n')[:12]:
        g = re.match(r'^#\s*([A-Z_]+):\s*(.*)$', line)
        if g: m[g.group(1)] = g.group(2).strip()
    return m

def blank_block_headings(text, maxlen):
    L = text.split('\n'); out = []
    for i, raw in enumerate(L):
        s = raw.strip()
        if not s or len(s) > maxlen or s.startswith('#'): continue
        if s in NOISE or PROV.search(s): continue
        if any(c in BAD for c in s): continue
        if not re.search(r'[㐀-鿿豈-﫿]', s): continue
        prev = L[i-1].strip() if i else ''
        nxt  = L[i+1].strip() if i+1 < len(L) else ''
        if prev or nxt: continue          # must sit alone in a blank-delimited block
        out.append(s)
    return out

def marker_headings(text, marker, maxlen):
    """◆-marked labels that run straight into body text: cut at the label."""
    out = []
    for raw in text.split('\n'):
        s = raw.strip()
        if not s.startswith(marker): continue
        s = s.lstrip(marker).strip()
        m = re.match(r'^([^　 ]{1,%d}?(?:（[^）]{1,6}）)?[^　 ]{0,6})(?=[，。“”「]|$)' % maxlen, s)
        lab = (m.group(1) if m else s)[:maxlen]
        if lab: out.append(lab)
    return out

def jkey(p):
    n = re.findall(r'(\d+)', os.path.basename(p))
    return (int(n[-1]) if n else 0, p)

#                work,        region, subdir,                              glob,        container, leaf,  mode,     maxlen, granularity
WORKS = [
 ('藝文類聚','中國','中國漢文/raw/唐朝/藝文類聚',            '*juan*.txt','部','子目','blank', 7,'category'),
 ('太平御覽','中國','中國漢文/raw/北宋/太平御覽',            '[0-9]*.txt','部','子目','blank', 9,'category'),
 ('冊府元龜','中國','中國漢文/raw/北宋/冊府元龜',            '*juan*.txt','部','門',  'blank', 8,'category'),
 ('雲笈七籤','中國','中國漢文/raw/北宋/雲笈七籤',            '*juan*.txt','部','篇目','blank', 9,'category'),
 ('事實類苑','中國','中國漢文/raw/南宋/事實類苑',            '*juan*.txt','門','子目','marker',10,'category'),
 ('四庫全書總目提要','中國','中國漢文/raw/清朝/四庫全書總目提要','*juan*.txt','部','類',  'blank', 8,'category'),
 ('古今事物考','中國','中國漢文/raw/清朝/古今事物考',        '*juan*.txt','類','子目','blank', 8,'category'),
 ('事物紀原','中國','中國漢文/raw/宋朝/事物紀原',            '*juan*.txt','部','子目','blank', 8,'category'),
 ('御定分類字錦','中國','中國漢文/raw/清朝/御定分類字錦',    '*juan*.txt','門','子目','blank', 8,'category'),
 # entry-granularity: real text, but the short lines are ENTRIES, not categories
 ('夜航船','中國','中國漢文/raw/明朝/夜航船',                '*juan*.txt','部','條目','blank', 8,'entry'),
 ('爾雅註疏','中國','中國漢文/raw/北宋/爾雅註疏',            '*juan*.txt','篇','條目','blank', 9,'entry'),
 ('通志','中國','中國漢文/raw/南宋/通志',                    '*juan*.txt','略','條目','blank', 8,'entry'),
 ('白氏六帖事類集','中國','中國漢文/raw/唐朝/白氏六帖事類集','*juan*.txt','部','門',  'blank', 8,'category'),
]

def run():
    os.makedirs(OUTDIR, exist_ok=True)
    vocab = load_vocab()
    rows, report = [], []
    for work, region, sub, pat, clevel, leaf, mode, maxlen, gran in WORKS:
        d = os.path.join(CORPUS, sub)
        files = sorted(glob.glob(os.path.join(d, pat)), key=jkey)
        voc = vocab.get(work, set())
        if not files:
            report.append(dict(work=work, files=0, rows=0, containers=0, leaves=0,
                               front=0, vocab=len(voc), granularity=gran,
                               note='no files found')); continue
        cur, cur_vol = '', None
        n_c = n_l = n_f = 0
        for f in files:
            text = read(f); m = meta_of(text)
            vol = m.get('PAGE_TITLE','').split('/')[-1] or os.path.splitext(os.path.basename(f))[0]
            src = m.get('URL','') or 'Fanya Hanwen corpus (Wikisource transcription)'
            hs  = marker_headings(text, '◆', maxlen) if mode == 'marker' else blank_block_headings(text, maxlen)
            order = 0
            for h in hs:
                order += 1
                if h == work or h.startswith(work):
                    continue                                  # running work title
                if h.endswith(FRONT) and len(h) <= 5:
                    n_f += 1
                    rows.append(dict(zip(COLS, ['', work, region, '卷前文字', '', '', order, vol, h, src,
                        '序跋凡例等卷前文字，非分類標目。'])));  continue
                if h in voc:                                   # attested container
                    cur, cur_vol = h, vol; n_c += 1
                    rows.append(dict(zip(COLS, ['', work, region, clevel, '', '', order, vol, h, src,
                        '見於《類書分類原目》既有著錄之上層標目。'])))
                else:
                    n_l += 1
                    if not cur:      note = '原文未見上層標目；不補造。'
                    elif cur_vol==vol: note = '上層標目見於同卷。'
                    else:            note = f'上層標目承自「{cur_vol}」卷，非本卷所載。'
                    if gran == 'entry':
                        note += ' 本書此層為條目，非分類標目，不得逕作類目使用。'
                    rows.append(dict(zip(COLS, ['', work, region, leaf, cur, h, order, vol, h, src, note])))
        report.append(dict(work=work, files=len(files), rows=n_c+n_l+n_f, containers=n_c,
                           leaves=n_l, front=n_f, vocab=len(voc), granularity=gran, note=''))
    for i, r in enumerate(rows, 1): r['Raw ID'] = f'CORP-{i:05d}'
    out = os.path.join(OUTDIR, 'leishu_原目_corpus.csv')
    with open(out, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=COLS); w.writeheader(); w.writerows(rows)
    json.dump(report, open(os.path.join(OUTDIR,'coverage.json'),'w',encoding='utf-8'),
              ensure_ascii=False, indent=2)
    print(f'WROTE {out}  rows={len(rows)}')
    print(f"{'work':18s}{'files':>6}{'rows':>7}{'cont':>6}{'leaf':>7}{'front':>6}{'vocab':>6}  granularity")
    for r in report:
        print(f"{r['work']:18s}{r['files']:6d}{r['rows']:7d}{r['containers']:6d}{r['leaves']:7d}"
              f"{r['front']:6d}{r['vocab']:6d}  {r['granularity']} {r['note']}")

COLS = ['Raw ID','Work','Region','Source level','Parent','Subdivision','Order',
        'Volume','Exact source label','Source','Notes']

if __name__ == '__main__':
    run()
