#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Two actions on fanya_category_grouping:

1. DROP the 四庫全書總目提要 rows from 類書分類原目. The 提要 are bibliographic
   abstracts, not a category system; the 四部 taxonomy already has its own sheet.

2. ADD 四庫著錄分類 — the 四庫 catalogue's own classification OF THE WORKS IT
   LISTS. 2,780 of 2,882 works carry 部/類, and 1,718 also carry 屬. Where the
   same work is held elsewhere in the corpus the classification propagates to
   that copy, so Browse by Category can place works that have no leishu
   attestation at all.

Titles are joined on the plain form: the 四庫 root suffixes them
「藝文類聚 (四庫全書本)」 while other roots hold 「藝文類聚」.
"""
import csv, collections, os, re
import openpyxl

REPO=os.path.expanduser('~/mnt/fanyahanwen-corpus')
BOOK=os.path.join(REPO,'fanya_category_grouping_2026-09-11.xlsx')
CORP=os.path.join(REPO,'corpus')
SUF=re.compile(r'\s*[（(](四庫全書本|四庫全書標點本|四庫全書薈要本)[）)]\s*$')
def plain(t): return SUF.sub('',(t or '').strip()).strip()

wb=openpyxl.load_workbook(BOOK)

# --- 1. drop the 提要 rows -------------------------------------------------
ws=wb['類書分類原目']; hdr=[c.value for c in ws[1]]; iW=hdr.index('Work')+1
drop=[r for r in range(2,ws.max_row+1)
      if str(ws.cell(r,iW).value or '').strip()=='四庫全書總目提要']
for r in sorted(drop,reverse=True): ws.delete_rows(r)
print(f'dropped 四庫全書總目提要 rows from 類書分類原目: {len(drop)}')

# --- 2. build 四庫著錄分類 --------------------------------------------------
pages=list(csv.DictReader(open(os.path.join(CORP,'四庫全書','index.csv'),encoding='utf-8-sig')))
works={}
for r in pages:
    t=(r.get('work_title') or '').strip()
    if not t: continue
    w=works.setdefault(t,dict(section='',c1='',c2='',author='',times='',pages=0,chars=0))
    w['pages']+=1
    try: w['chars']+=int(float(r.get('char_count_clean') or 0))
    except Exception: pass
    for k,src in (('section','section'),('c1','category1'),('c2','category2'),
                  ('author','author'),('times','times')):
        v=(r.get(src) or '').strip()
        if v and not w[k]: w[k]=v

v2=list(csv.DictReader(open(os.path.join(CORP,'index_corpus_v2.csv'),encoding='utf-8-sig')))
elsewhere=collections.defaultdict(list)
for r in v2:
    if r['corpus_root']=='四庫全書': continue
    elsewhere[plain(r['work'])].append(f"{r['corpus_root']}/{r['grouping']}".rstrip('/'))

name='四庫著錄分類'
if name in wb.sheetnames: del wb[name]
sh=wb.create_sheet(name)
sh.append(['書名','部','類','屬','著者','時代','四庫本卷數','四庫本字數',
           '四庫本題名','語料庫他處收藏','分類完整度','備註'])
n_full=n_prop=0
for t,w in sorted(works.items(), key=lambda kv:(kv[1]['section'],kv[1]['c1'],kv[0])):
    p=plain(t); other=sorted(set(elsewhere.get(p,[])))
    if other: n_prop+=1
    lvl=('部類屬' if w['c2'] else ('部類' if w['c1'] else ('部' if w['section'] else '未著錄')))
    if w['c1']: n_full+=1
    sh.append([p, w['section'], w['c1'], w['c2'], w['author'], w['times'],
               w['pages'], w['chars'], t, '；'.join(other), lvl,
               ('同書亦見於語料庫他處，可據此賦類。' if other else '')])
for col,wd in zip('ABCDEFGHIJKL',(24,8,14,14,14,10,11,12,28,34,10,30)):
    sh.column_dimensions[col].width=wd
wb.save(BOOK)
print(f'新增工作表「{name}」: {sh.max_row-1} 部書')
print(f'  有部類者          : {n_full}')
print(f'  可外推至他處收藏者 : {n_prop}')
sec=collections.Counter(w['section'] for w in works.values() if w['section'])
print('  部別:',dict(sec))
