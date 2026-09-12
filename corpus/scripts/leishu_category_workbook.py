# -*- coding: utf-8 -*-
"""Add the category-context join to the grouping workbook.

  類目徵引   one row per (leishu, category path): what that category actually quotes
  四庫著錄分類 gains 類書類目 — the leishu category labels a catalogued work sits under
"""
import csv, collections, openpyxl, os
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

WB='fanya_category_grouping_2026-09-11.xlsx'
SRC='corpus/index_leishu_徵引_分類.csv'

rows=list(csv.DictReader(open(SRC, encoding='utf-8-sig')))
by=collections.defaultdict(list); leaves=collections.defaultdict(collections.Counter)
gen={}
for r in rows:
    n=int(r['citations'])
    by[(r['leishu'], r['category_path'])].append((r['cited_work'], n))
    gen[(r['leishu'], r['category_path'])]=r['genmoku']=='1'
    leaves[r['cited_work']][r['category_leaf']]+=n

wb=openpyxl.load_workbook(WB)
if '類目徵引' in wb.sheetnames: del wb['類目徵引']
ws=wb.create_sheet('類目徵引')
hdr=['類書','類目路徑','類目','原目確認','徵引次數','所徵引書數','所徵引之書']
ws.append(hdr)
for c in range(1, len(hdr)+1): ws.cell(1, c).font=Font(bold=True)
out=[]
for (leishu, path), items in by.items():
    items.sort(key=lambda kv:(-kv[1], kv[0]))
    out.append([leishu, path, path.split(' > ')[-1],
                '是' if gen[(leishu, path)] else '未詳',
                sum(n for _, n in items), len(items),
                '、'.join(f'{w}({n})' for w, n in items[:12])])
out.sort(key=lambda r:(r[0], -r[4], r[1]))
for r in out: ws.append(r)
for col, w in zip('ABCDEFG', (16, 34, 14, 10, 10, 12, 90)):
    ws.column_dimensions[col].width=w
ws.freeze_panes='A2'
ws.auto_filter.ref=f'A1:{get_column_letter(len(hdr))}{ws.max_row}'

# 四庫著錄分類 gains the leishu category labels for each catalogued work
sk=wb['四庫著錄分類']
head=[c.value for c in sk[1]]
if '類書類目' not in head:
    col=len(head)+1
    sk.cell(1, col, '類書類目').font=Font(bold=True)
    sk.column_dimensions[get_column_letter(col)].width=40
else:
    col=head.index('類書類目')+1
n=0
for r in range(2, sk.max_row+1):
    title=sk.cell(r, 1).value
    lv=leaves.get(title)
    if not lv: continue
    sk.cell(r, col, '、'.join(f'{k}({v})' for k, v in lv.most_common(8)))
    n+=1
wb.save(WB)
print(f'類目徵引: {len(out):,} rows')
print(f'四庫著錄分類 類書類目 filled: {n} works')
